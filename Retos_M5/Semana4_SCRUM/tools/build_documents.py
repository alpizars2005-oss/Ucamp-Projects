"""Completa las plantillas recuperadas y compone los anexos, sin fingir eventos.

Requiere python-docx para DOCX. --pdf requiere LibreOffice y pypdf.
La aplicación no necesita estas dependencias de documentación.
"""
from copy import deepcopy
from datetime import datetime, timezone
import io
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import argparse
import re
import zipfile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image

B = Path(__file__).resolve().parents[1]
D = json.loads((B / 'tools/report_data.json').read_text(encoding='utf-8'))
OUT = B / 'documentos'
PINK = 'D9268B'
INK = '272C30'


def paragraph(doc, text='', bold=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    return p


def title(doc, text, label=None):
    if label: paragraph(doc, label, True, 9)
    doc.add_heading(text, 1)


def page(doc, text, label='ANEXO · SEMANA 4'):
    doc.add_page_break()
    title(doc, text, label)


def simple_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    if 'Table Grid' not in doc.styles:
        style = doc.styles.add_style('Table Grid', WD_STYLE_TYPE.TABLE)
        pr = OxmlElement('w:tblPr'); borders = OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            item = OxmlElement('w:'+edge); item.set(qn('w:val'),'single'); item.set(qn('w:sz'),'4'); item.set(qn('w:color'),'D5D0D2'); borders.append(item)
        pr.append(borders); style._element.append(pr)
    t.style = 'Table Grid'
    for i, h in enumerate(headers): t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row): cells[i].text = str(value)
    for ridx, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        for cidx, c in enumerate(row.cells):
            if widths: c.width = Inches(widths[cidx])
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.bold = ridx == 0
            if ridx == 0:
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'F7DEEC'); c._tc.get_or_add_tcPr().append(shd)
    hdr = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(hdr)
    return t


def screenshot(doc, filename, caption):
    # JPEG copy only reduces document size; original PNG evidence is retained intact.
    image = Image.open(B / 'evidencias' / filename).convert('RGB')
    buffer = io.BytesIO(); image.save(buffer, format='JPEG', quality=88)
    buffer.seek(0)
    shape = doc.add_picture(buffer, width=Inches(6.85))
    shape._inline.docPr.set('descr', caption)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, caption, False, 9)


def fill_template():
    doc = Document(B / 'tools/templates/Semana4_base.docx')
    t = doc.tables[0]
    # Original Google export includes oversized/decimal row heights. Remove only
    # layout constraints; all official fields and their order remain present.
    for row in t.rows:
        for e in list(row._tr.get_or_add_trPr()):
            if e.tag in (qn('w:trHeight'), qn('w:cantSplit')): e.getparent().remove(e)
    def cell(r, text, col=0): t.rows[r].cells[col].text = text
    cell(0,'Mochi Monchi · ejercicio individual',1)
    cell(1,'Angel Alfredo · responsabilidad simulada de SM',1)
    cell(2,'Angel Alfredo · responsabilidad simulada de PO',1)
    cell(3,'Angel Alfredo · responsabilidad simulada de Developer',1)
    cell(1,'1 semana simulada\n5 jornadas virtuales\nJ1–J5\n\nPreparación:\n30/09/2026',2)
    for r,(sid,name,criterion,evidence) in zip(range(5,13),D['stories'][:8]):
        cell(r,f'{sid} · {name}. {criterion}')
    cell(13,'US-09 · Resumen de pedidos y total del día. US-10 · Valoración de 1 a 5 y comentario opcional.\nRelease 2: cancelación, pagos/tickets, umbrales, ajustes, permisos y exportación. Futuro: nube y multi-sucursal.')
    for i in range(3): cell(15+i,f'{2*i+1}. {D["dor"][2*i]}\n{2*i+2}. {D["dor"][2*i+1]}')
    for i in range(3): cell(19+i,f'{2*i+1}. {D["dod"][2*i]}\n{2*i+2}. {D["dod"][2*i+1]}')
    cell(23,'US-01 a US-10, seleccionadas del MVP de Semana 3. Sin puntos ni velocidad histórica inventados.')
    cell(24,'TABLERO DE LA SIMULACIÓN\nPlanning: por hacer US-01–US-10.\nJornada intermedia didáctica: terminadas US-01–02; en revisión US-03–05; en curso US-06–08; pendientes US-09–10.\nCierre verificado: US-01–US-10 terminadas en el alcance académico local.\nNo es una captura ni un historial de Jira/Miro.')
    cell(26,'Registrar un pedido una sola vez, revisar su total, confirmarlo y actualizar ingredientes y empaques sin doble captura. Si falta un insumo, rechazar sin cambios. Al cierre, consultar pedidos del día y registrar feedback.')
    cell(28,'Técnica: Empezar / Dejar / Continuar.\nHallazgos: la integridad de los datos necesita probar errores; una nota no debe cambiar la receta; la validación comercial sigue pendiente.\nMejoras: prueba con operador real, cancelación con devolución de stock y planificación de respaldo/permisos. Responsables y comprobaciones: anexo de Review y documento de observaciones.')
    cell(30,'\n\n'.join(q+'\n'+a for q,a in D['lessons']))
    for row in t.rows:
        seen=set()
        for c in row.cells:
            if id(c._tc) in seen: continue
            seen.add(id(c._tc))
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.name = 'Arial'; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor.from_string(INK)
    for r in [4,14,18,22,25,27,29]:
        for p in t.rows[r].cells[0].paragraphs:
            for run in p.runs:
                run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    for row, col in [(0,0),(0,2),(1,0),(2,0),(3,0)]:
        for p in t.rows[row].cells[col].paragraphs:
            for run in p.runs: run.font.color.rgb = RGBColor(255,255,255); run.bold = True
    tbls=[]
    for start,end in [(0,14),(14,27),(27,31)]:
        copy = deepcopy(t._tbl)
        for i,tr in enumerate(list(copy.findall(qn('w:tr')))):
            if not start <= i < end: copy.remove(tr)
        tbls.append(copy)
    # Reuse the official table, headers, section and styles, with explicit page divisions.
    body=doc._element.body
    for item in list(body):
        if item.tag != qn('w:sectPr'): body.remove(item)
    section=doc.sections[0]
    section.top_margin=Inches(.90); section.bottom_margin=Inches(.62)
    section.left_margin=Inches(.65); section.right_margin=Inches(.65)
    section.header_distance=Inches(.2); section.footer_distance=Inches(.2)
    for name in ['normal']:
        style=doc.styles[name]
        style.font.name='Arial'; style.font.size=Pt(10)
        style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.line_spacing=1.12
    for name in ['Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor.from_string(PINK)
        style.font.size=Pt(18 if name=='Heading 1' else 12)
    doc.core_properties.author='Angel Alfredo'
    doc.core_properties.last_modified_by='Angel Alfredo'
    doc.core_properties.title='Mochi Monchi · Simulación de Sprint e Incremento'
    # Preserve branding while providing an unambiguous page identifier.
    section.header.paragraphs[0].text='UCAMP · Proyectos Ágiles con SCRUM · Mochi Monchi'
    section.header.paragraphs[0].runs[0].font.size=Pt(8)
    section.footer.paragraphs[0].text='Angel Alfredo · Semana 4 · Simulación académica individual'
    section.footer.paragraphs[0].runs[0].font.size=Pt(8)
    section.footer.paragraphs[0].add_run(' · Página ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); section.footer.paragraphs[0]._p.append(field)
    title(doc,'Observaciones del Sprint','PLANTILLA OFICIAL COMPLETADA · SEMANA 4')
    paragraph(doc,'Mochi Monchi — sistema digital de pedidos e inventario',True,11)
    paragraph(doc,'Simulación individual: una semana representada en cinco jornadas virtuales, no cinco días transcurridos. No se inventan integrantes ni reuniones reales. Los resultados de software sí proceden de pruebas ejecutadas.',False,9)
    body.insert(len(body)-1,tbls[0])
    page(doc,'Acuerdos y objetivo del Sprint','PLANTILLA OFICIAL · CONTINUACIÓN')
    body.insert(len(body)-1,tbls[1])
    page(doc,'Retrospectiva y aprendizajes','PLANTILLA OFICIAL · CONTINUACIÓN')
    paragraph(doc,'Las reflexiones se proponen para revisión del estudiante antes de la entrega. La simulación no sustituye la experiencia de colaborar con un equipo real.',False,9)
    body.insert(len(body)-1,tbls[2])
    return doc


def build():
    OUT.mkdir(exist_ok=True)
    shutil.copyfile(B/'tools/templates/Semana2_completada.docx',OUT/'UCAMP_SCRUM_Semana2_Mochi_Monchi.docx')
    doc=fill_template()
    page(doc,'Sprint Planning')
    paragraph(doc,'Guion: inicio de J1 · Timebox propuesto: 45 minutos, no tiempo medido.',True,10)
    paragraph(doc,'Por qué',True)
    paragraph(doc,'Comprobar si una sola captura puede alimentar pedidos, recetas e inventario. La hipótesis de ahorrar tiempo y reducir errores de operación todavía requiere una prueba con usuarios reales.')
    paragraph(doc,'Qué',True)
    paragraph(doc,'Se seleccionan las diez historias del MVP de Semana 3. Se limitan los datos a tres productos y recetas de demostración. Quedan fuera pagos, cancelación, reposición, permisos, exportación y nube. No se utiliza una velocidad histórica ni capacidad de equipo inventadas.')
    paragraph(doc,'Cómo',True)
    paragraph(doc,'Construir el flujo pedido–inventario antes del cierre: interfaz de navegador, API Python y base SQLite local. El servidor calcula importes en centavos y suma ingredientes compartidos. Una transacción guarda la venta y descuenta existencias; un identificador único permite reintentar sin duplicar.')
    simple_table(doc,['Acuerdo','Motivo'],[
        ('Probar fallos como parte de cada historia','Una pantalla correcta no demuestra integridad del inventario.'),
        ('Distinguir nota y receta','La nota sólo informa; no cambia cantidades ni ingredientes.'),
        ('Una base compartida, no copias separadas','Los navegadores consultan el mismo servidor; no hay sincronización en la nube.'),
        ('No usar datos o clientes reales en la evidencia','Precios, recetas, pedidos y valoraciones son sintéticos.')])
    paragraph(doc,'Salida de Planning: objetivo, Sprint Backlog US-01–US-10, criterios de aceptación, DoR/DoD y plan de pruebas. Los acuerdos DoR son una práctica del ejercicio, no un requisito adicional del marco Scrum.',False,10)
    paragraph(doc,'El Sprint funciona como contenedor del trabajo y de los cuatro eventos documentados. La duración propuesta es una semana; las jornadas siguientes representan su simulación, no un calendario observado.',False,10)
    page(doc,'Daily Scrum · cinco jornadas virtuales')
    paragraph(doc,'Timebox propuesto: diez minutos por jornada. Se inspecciona el avance hacia el objetivo y se adapta el plan. No es un reporte a un jefe ni un registro de reuniones reales.',False,10)
    for day,name,observed,decision,limit in D['dailies']:
        doc.add_heading(day+' · '+name,2)
        paragraph(doc,observed+' '+decision,False,9.5)
        paragraph(doc,'Impedimento o límite: '+limit,False,9)
    page(doc,'Sprint Review · resultado e incremento')
    paragraph(doc,'Guion: fin de J5 · Timebox propuesto: 25 minutos. Se representan las perspectivas de PO, desarrollo y operador; no hubo aprobación de un cliente real.',False,10)
    simple_table(doc,['Demostración','Resultado técnico observado'],[
        ('2 fresa + 1 matcha','Total previo: $110.00. Se guarda un pedido y se descuentan receta y empaques.'),
        ('Segundo pedido: 10 fresa','Total acumulado: 2 pedidos y $460.00. Masa 80 g, crema 40 g y fresa 20 g generan tres alertas.'),
        ('Intentar 3 matcha sin stock suficiente','Se rechaza. El estado completo de pedidos e inventario no cambia.'),
        ('Cierre y feedback','Se guarda 4/5 y un comentario sintético; no representa satisfacción de un cliente.'),
        ('Respuesta perdida y reintento','Una sola venta y un solo descuento; el mismo envío no se procesa dos veces.')])
    paragraph(doc,'Aceptación de la simulación',True)
    paragraph(doc,'US-01 a US-10 cumplen los criterios dentro del alcance académico local. La Review inspecciona el incremento y orienta el backlog; no sustituye la DoD. No se concluye que el sistema esté listo para uso productivo o que ya reduzca tiempos reales.')
    paragraph(doc,'Ajuste del backlog',True)
    paragraph(doc,'Para el siguiente Sprint se prioriza una prueba con operador real y cancelación con devolución de stock. Antes de datos operativos se deben definir respaldo/restauración y permisos. El feedback sintético sirve para verificar que se guarda, no para tomarlo como evidencia comercial.')
    paragraph(doc,'Retrospectiva: acciones comprobables',True)
    simple_table(doc,['Acción','Responsable y señal de cumplimiento'],[
        ('Empezar una prueba con operador','PO: cinco pedidos, registro de tiempo/correcciones y comparación con una línea base manual.'),
        ('Empezar cancelación segura','Developer: stock restituido exactamente, doble cancelación rechazada y cierre corregido.'),
        ('Dejar supuestos en notas','PO: mantener advertencia; no aceptar sustituciones sin definir su receta.'),
        ('Continuar las pruebas de errores','SM/Developer: conservar pruebas de concurrencia, idempotencia y atomicidad sin regresiones.')])
    page(doc,'Evidencia 1 · revisar antes de confirmar')
    screenshot(doc,'01_pedido_revision.png','Captura real de la interfaz durante la comprobación: dos mochi de fresa y uno de matcha. Datos de demostración, no una venta real.')
    simple_table(doc,['Elemento','Comprobación'],[
        ('US-01 / US-03 / US-04','Catálogo con precios, carrito de varios productos y nota breve.'),
        ('US-05','2 × $35.00 + 1 × $40.00 = $110.00 MXN.'),
        ('Límite','La nota no altera la receta. El total definitivo se recalcula en el servidor.')])
    page(doc,'Evidencia 2 · inventario después de vender')
    screenshot(doc,'03_inventario_alertas.png','Captura real después del segundo pedido sintético. Las tres alertas corresponden a masa, crema y fresa, sin existencias negativas.')
    paragraph(doc,'Después del primer pedido: masa 480 g; crema 240 g; fresa 170 g; mango 200 g; matcha 65 g; empaques 17. Después del segundo: 80, 40, 20, 200, 65 y 7, respectivamente. El consumo se deriva de recetas fijas de prueba.',False,10)
    paragraph(doc,'US-07 y US-08: una confirmación actualiza pedido e inventario en la misma transacción. El intento posterior de tres matcha es rechazado sin modificar ese estado; captura adicional: evidencias/06_rechazo_sin_stock.png.',False,10)
    page(doc,'Evidencia 3 · cierre y feedback')
    screenshot(doc,'04_cierre_y_feedback.png','Captura real del cierre: dos pedidos y $460.00 de demostración. La valoración 4/5 y su comentario son sintéticos y están identificados como prueba.')
    paragraph(doc,'US-09 y US-10: resumen del día y captura de feedback persistente. El día se determina por la fecha local del servidor. La valoración no demuestra satisfacción ni validación del mercado.',False,10)
    paragraph(doc,'Otras capturas entregadas: 02_pedido_confirmado.png, 05_vista_movil.png y 06_rechazo_sin_stock.png. La vista móvil usa un contexto de navegador de 390 píxeles; no acredita una prueba en un teléfono físico.',False,10)
    page(doc,'Verificación y consolidado final')
    result=json.loads((B/'evidencias/verificacion_navegador.json').read_text())
    paragraph(doc,'44 pruebas de lógica/API y nueve comprobaciones de interfaz aprobadas. Seis capturas y cero errores de JavaScript observados.',True,11)
    paragraph(doc,'Entorno registrado: '+result['environment']+'.',False,9)
    paragraph(doc,'Las pruebas de interfaz usan datos temporales. El modo de puente HTTP, cuando está indicado, no verifica navegación nativa, almacenamiento nativo ni políticas del navegador. Dos contextos no equivalen a dos equipos físicos. El registro JSON conserva el modo y los resultados exactos.',False,9)
    paragraph(doc,'Comando de pruebas desde la raíz del repositorio:',True,9)
    paragraph(doc,'python -m unittest discover -s Retos_M5/Semana4_SCRUM/tests -v',False,8)
    simple_table(doc,['Rúbrica','Puntos','Evidencia'],[
        ('Eventos Scrum','3','Planning, cinco Daily virtuales, Review, Retrospective y duración del Sprint.'),
        ('Incremento funcional','3','Código, criterios US-01–US-10, pruebas y capturas.'),
        ('Reflexiones','2','Aprendizajes y acciones de mejora; revisión personal pendiente antes de Community.'),
        ('Ambas plantillas','2','Semana 2 recuperada y esta Semana 4 completada con anexos.')])
    paragraph(doc,'La tabla muestra cobertura de la entrega, no una calificación otorgada.',False,9)
    paragraph(doc,'Enlaces y archivos',True,10)
    paragraph(doc,D['repo_url'],False,8)
    paragraph(doc,'documentos/: ambas plantillas DOCX/PDF y UCAMP_SCRUM_Consolidado.pdf. evidencias/: seis capturas, pruebas.txt y verificacion_navegador.json. ENTREGA_COMMUNITY.txt: texto de acompañamiento. Publicar en GitHub no equivale a enviar a Community.',False,9)
    paragraph(doc,'Referencias',True,10)
    for text in [
        'UCAMP. M1S4 Observaciones del Sprint. Plantilla oficial facilitada en la consigna.',
        'https://docs.google.com/document/d/1d07cLEUY9ldEh5usOPw_lTNAoIRd36x3mDNNMCkRFlU/edit',
        'Schwaber, K. y Sutherland, J. (2020). The Scrum Guide. https://scrumguides.org/scrum-guide.html',
        'Python Software Foundation. sqlite3 y http.server. https://docs.python.org/3/library/sqlite3.html · https://docs.python.org/3/library/http.server.html',
        'Continuidad del proyecto: Product_Vision_Board.md (Semana 2) y User_Story_Mapping_MVP.md (Semana 3), en este repositorio.'
    ]: paragraph(doc,text,False,8)
    doc.save(OUT/'UCAMP_SCRUM_Semana4_Mochi_Monchi.docx')
    # No distributable DOCX may include raw font binaries.
    for path in OUT.glob('*.docx'):
        with zipfile.ZipFile(path) as z:
            assert not any(n.startswith('word/fonts/') for n in z.namelist())
    print('DOCX:',OUT)


def export_pdf():
    from pypdf import PdfWriter
    lo=shutil.which('libreoffice') or shutil.which('soffice')
    if not lo: raise SystemExit('Instala LibreOffice para exportar PDF.')
    with tempfile.TemporaryDirectory() as profile:
        for name in ['UCAMP_SCRUM_Semana2_Mochi_Monchi','UCAMP_SCRUM_Semana4_Mochi_Monchi']:
            subprocess.run([lo,f'-env:UserInstallation={Path(profile).as_uri()}', '--headless','--convert-to','pdf','--outdir',str(OUT),str(OUT/(name+'.docx'))],check=True,timeout=120)
            if not (OUT/(name+'.pdf')).exists():raise RuntimeError('La conversión no produjo PDF: '+name)
    writer=PdfWriter()
    for name in ['UCAMP_SCRUM_Semana2_Mochi_Monchi','UCAMP_SCRUM_Semana4_Mochi_Monchi']:
        writer.append(OUT/(name+'.pdf'),outline_item=name.replace('_',' '))
    writer.add_metadata({'/Title':'Mochi Monchi · Consolidado Scrum · Semanas 2 y 4','/Author':'Angel Alfredo'})
    with (OUT/'UCAMP_SCRUM_Consolidado.pdf').open('wb') as f:writer.write(f)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pdf',action='store_true')
    args=parser.parse_args()
    build()
    if args.pdf:export_pdf()
